"""Đọc / ghi file Excel dùng chung cho mọi import và export.

Đọc: nhận diện cột theo TÊN — không phân biệt hoa thường, dấu, gạch dưới — chứ không theo vị trí,
vì BTC hay chèn thêm cột ghi chú. Chỉ nhận .xlsx thật: kiểm tra chữ ký file, không tin đuôi tên.

Ghi: mọi ô chữ bị ép kiểu CHUỖI. openpyxl tự coi chuỗi bắt đầu bằng "=" là công thức, nên một họ tên
CBNV tự khai kiểu `=HYPERLINK(...)` sẽ chạy khi BTC mở file (Excel/CSV injection). Ép kiểu chuỗi
cũng giữ được số 0 đầu của số điện thoại.
"""

import io
import re
import unicodedata
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.config import settings
from app.core.exceptions import AppError

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# File .xlsx là một kho zip: 4 byte đầu luôn là "PK\x03\x04".
XLSX_SIGNATURE = b"PK\x03\x04"
MAX_ROWS = 2000
TRUTHY = {"x", "1", "co", "true", "yes", "y"}

_HEADER_FILL = PatternFill("solid", fgColor="E2E8F0")
_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


# --- Đọc ---


def normalize(text: object) -> str:
    """'  Mã  NV ' -> 'ma nv'. Bỏ dấu, gộp khoảng trắng, gạch dưới coi như khoảng trắng."""
    value = str(text or "").strip().lower().replace("đ", "d")
    value = unicodedata.normalize("NFD", value)
    value = "".join(char for char in value if unicodedata.category(char) != "Mn")
    return " ".join(value.replace("_", " ").split())


def build_aliases(raw: dict[str, set[str]]) -> dict[str, set[str]]:
    return {name: {normalize(alias) for alias in aliases} for name, aliases in raw.items()}


def cell_text(value: object) -> str:
    """Giá trị ô -> chữ. Số 801.0 -> '801'; ô ngày -> '2021-03-01'."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        if (value.hour, value.minute, value.second) == (0, 0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_rows(
    content: bytes,
    *,
    aliases: dict[str, set[str]],
    validate_header: Callable[[dict[str, int], list[str]], None],
    max_rows: int = MAX_ROWS,
) -> list[dict[str, Any]]:
    """Đọc sheet đầu tiên thành các dòng `{field: text, "row": số dòng Excel}`.

    `validate_header(mapping, cells)` ném AppError nếu thiếu cột bắt buộc.
    """
    if not content:
        raise AppError("File rỗng.", code="EMPTY_FILE")

    max_bytes = settings.MAX_UPLOAD_MB * 1024 * 1024
    if len(content) > max_bytes:
        raise AppError(
            f"File vượt quá {settings.MAX_UPLOAD_MB}MB.",
            code="FILE_TOO_LARGE",
            details={"max_mb": settings.MAX_UPLOAD_MB, "actual_bytes": len(content)},
        )
    if not content.startswith(XLSX_SIGNATURE):
        raise AppError("Chỉ nhận file Excel định dạng .xlsx.", code="UNSUPPORTED_FILE_TYPE")

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # file zip hỏng hoặc không phải workbook
        raise AppError(
            "Không đọc được file Excel. Mở lại bằng Excel rồi lưu dạng .xlsx.",
            code="UNSUPPORTED_FILE_TYPE",
        ) from exc

    header: dict[str, int] | None = None
    parsed: list[dict[str, Any]] = []
    try:
        sheet = workbook.worksheets[0]
        for number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            cells = [cell_text(value) for value in values]
            if not any(cells):
                continue
            if header is None:
                header = map_header(cells, aliases)
                validate_header(header, cells)
                continue

            record = {
                name: (cells[index] if index < len(cells) else "") for name, index in header.items()
            }
            record["row"] = number
            parsed.append(record)
            if len(parsed) > max_rows:
                raise AppError(
                    f"File có hơn {max_rows} dòng. Chia nhỏ file rồi import từng phần.",
                    code="TOO_MANY_ROWS",
                )
    finally:
        workbook.close()

    if header is None:
        raise AppError("File không có dòng tiêu đề.", code="MISSING_COLUMNS")
    return parsed


def map_header(cells: list[str], aliases: dict[str, set[str]]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for index, cell in enumerate(cells):
        key = normalize(cell)
        for name, names in aliases.items():
            if key in names and name not in mapping:
                mapping[name] = index
    return mapping


def parse_date(text: str) -> str | None:
    """'2020-01-31', '2020-01-31 08:00', '31/01/2020', '31-1-2020' -> '2020-01-31'. Sai -> None."""
    text = text.strip()
    if not text:
        return None
    iso = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:$|[ T])", text)
    if iso:
        year, month, day = iso.groups()
    else:
        vietnamese = re.match(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})$", text)
        if not vietnamese:
            return None
        day, month, year = vietnamese.groups()
    try:
        return date(int(year), int(month), int(day)).isoformat()
    except ValueError:
        return None


# --- Ghi ---


@dataclass
class Sheet:
    title: str
    headers: list[str]
    rows: list[list[Any]] = field(default_factory=list)


def build_workbook(sheets: list[Sheet]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_titles: set[str] = set()

    for sheet in sheets:
        worksheet = workbook.create_sheet(_sheet_title(sheet.title, used_titles))
        for column, header in enumerate(sheet.headers, start=1):
            cell = worksheet.cell(row=1, column=column, value=header)
            cell.font = Font(bold=True)
            cell.fill = _HEADER_FILL
            cell.alignment = Alignment(vertical="center")
        for row_number, values in enumerate(sheet.rows, start=2):
            for column, value in enumerate(values, start=1):
                _write_cell(worksheet, row_number, column, value)

        worksheet.freeze_panes = "A2"
        if sheet.rows:
            worksheet.auto_filter.ref = worksheet.dimensions
        for column, header in enumerate(sheet.headers, start=1):
            worksheet.column_dimensions[get_column_letter(column)].width = _width(
                header, sheet.rows, column - 1
            )

    if not workbook.worksheets:
        workbook.create_sheet("Trống")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def safe_filename(stem: str) -> str:
    """Tên file ASCII an toàn cho header Content-Disposition: 'Đăng ký TB2026' -> 'dang-ky-tb2026'."""
    text = normalize(stem)
    return re.sub(r"[^a-z0-9]+", "-", text).strip("-") or "export"


def _write_cell(worksheet, row: int, column: int, value: Any) -> None:
    if value is None or value == "":
        return
    if isinstance(value, bool):
        value = "Có" if value else "Không"
    cell = worksheet.cell(row=row, column=column)
    if isinstance(value, int | float):
        cell.value = value
        return
    cell.value = str(value)
    # Luôn là chuỗi: không để "=..." thành công thức, không để "0912..." mất số 0.
    cell.data_type = "s"


def _sheet_title(title: str, used: set[str]) -> str:
    base = _INVALID_SHEET_CHARS.sub(" ", title).strip()[:31] or "Sheet"
    candidate, suffix = base, 2
    while candidate.lower() in used:
        tail = f" ({suffix})"
        candidate = base[: 31 - len(tail)] + tail
        suffix += 1
    used.add(candidate.lower())
    return candidate


def _width(header: str, rows: list[list[Any]], index: int) -> int:
    longest = len(header)
    for values in rows[:200]:
        if index < len(values) and values[index] is not None:
            longest = max(longest, len(str(values[index])))
    return max(8, min(longest + 2, 60))
