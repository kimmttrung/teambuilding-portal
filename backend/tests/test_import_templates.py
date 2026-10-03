"""File mẫu tải được, có ví dụ/hướng dẫn và đúng cấu trúc mà importer đọc."""

from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models import Event
from app.models.accommodation import Hotel, Room
from app.models.enums import EventStatus, UserRole
from app.services import room_import_service, user_import_service
from app.services.excel import read_rows


@pytest.fixture
def template_world(db, make_user, auth_headers):
    make_user(email="btc@company.vn", role=UserRole.ADMIN)
    make_user(email="nv@company.vn")
    event = Event(
        code="TEMPLATE",
        name="Mẫu Excel",
        start_date="2026-10-15",
        end_date="2026-10-17",
        status=EventStatus.REGISTRATION_OPEN,
        is_active=True,
    )
    db.add(event)
    db.flush()
    hotel = Hotel(event_id=event.id, name="Khách sạn Phú Quốc")
    db.add(hotel)
    db.flush()
    db.add(Room(hotel_id=hotel.id, room_number="0012", capacity=3, gender_policy="any"))
    db.commit()
    return auth_headers("btc@company.vn"), auth_headers("nv@company.vn")


@pytest.mark.parametrize(
    "url", ["/api/v1/admin/users/import-template", "/api/v1/rooms/import-template"]
)
def test_templates_have_examples_and_importable_headers(client, template_world, url):
    admin, employee = template_world
    assert client.get(url, headers=employee).status_code == 403
    response = client.get(url, headers=admin)
    assert response.status_code == 200, response.text
    assert "attachment" in response.headers["content-disposition"]
    book = load_workbook(BytesIO(response.content))
    assert "Hướng dẫn" in book.sheetnames
    assert book.worksheets[0].max_row >= 3
    assert book["Hướng dẫn"].max_row > 3
    if "/users/" in url:
        rows = read_rows(
            response.content,
            aliases=user_import_service.COLUMN_ALIASES,
            validate_header=user_import_service._check_header,
        )
        assert rows[0]["phone"] == "0912345678"
        assert rows[0]["id_card_number"] == "001095012345"
        assert rows[0]["role"] == "CBNV"
        assert all(rows[0][name] for name in user_import_service.COLUMN_ALIASES)
    else:
        rows = room_import_service.parse_rows(response.content)
        assert rows[0]["room_number"] == "0012"
        assert rows[0]["hotel"] == "Khách sạn Phú Quốc"
        assert rows[0]["is_room_captain"] == "x"
        assert "Danh mục phòng" in book.sheetnames
    book.close()
