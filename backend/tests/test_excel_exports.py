"""Kiểm thử xuất Excel cho BTC (bước 21, docs/09-security.md §7.4)."""

import json
import zipfile
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.accommodation import Hotel, Room, RoomAssignment
from app.models.audit import AuditLog
from app.models.enums import (
    AssignmentMode,
    EventStatus,
    FlightDirection,
    Gender,
    RegistrationStatus,
    RoomGenderPolicy,
    UserRole,
)
from app.models.event import Event
from app.models.flight import Flight, FlightAssignment, Shift
from app.models.org import Team
from app.models.registration import Registration, RegistrationBusNeed
from app.models.transportation import Bus, BusAssignment, PickupPoint, TripLeg

NOW = "2026-09-12T04:00:00+00:00"
EVIL_NAME = '=HYPERLINK("http://evil.example","Bấm vào")'
EXPORTS = [
    "/api/v1/admin/users/export",
    "/api/v1/registrations/export",
    "/api/v1/flights/export",
    "/api/v1/buses/export",
    "/api/v1/rooms/export",
]


@pytest.fixture
def world(db: Session, make_user) -> dict:
    """Người đi đã được xếp bay/xe/phòng; bạn cùng phòng chưa có chuyến và chưa có xe; 1 người chưa đăng ký."""
    event = Event(
        code="TB2026", name="Team Building 2026", start_date="2026-10-15", end_date="2026-10-17",
        status=EventStatus.REGISTRATION_CLOSED, terms_version="v1", is_active=True,
    )
    team = Team(code="IT", name="Công nghệ")
    db.add_all([event, team])
    db.flush()
    shift = Shift(event_id=event.id, code="CA1", name="Ca 1", display_order=1)
    db.add(shift)
    db.flush()
    outbound = Flight(
        event_id=event.id, flight_code="VN1234", airline="Vietnam Airlines",
        direction=FlightDirection.OUTBOUND, shift_id=shift.id, departure_airport="HAN",
        arrival_airport="PQC", departure_time="2026-10-15T06:30:00+00:00",
        arrival_time="2026-10-15T08:40:00+00:00", capacity=60,
    )
    leg = TripLeg(
        event_id=event.id, code="CITY_TO_AIRPORT", name="HN → Sân bay",
        direction=FlightDirection.OUTBOUND, leg_date="2026-10-15", is_airport_linked=True,
        display_order=1,
    )
    db.add_all([outbound, leg])
    db.flush()
    pickup = PickupPoint(event_id=event.id, trip_leg_id=leg.id, name="Toà nhà Keangnam")
    db.add(pickup)
    db.flush()

    make_user(email="btc@company.vn", role=UserRole.ADMIN, full_name="Ban Tổ Chức", employee_code="BTC001")
    traveller = make_user(
        email="nv@company.vn", full_name=EVIL_NAME, employee_code="NV001", phone="0911000111",
        team_id=team.id, gender=Gender.MALE, date_of_birth="1995-01-01",
        id_card_number="001095012345", health_note="Dị ứng hải sản",
    )
    mate = make_user(
        email="mate@company.vn", full_name="Trần Văn Cùng", employee_code="NV002",
        team_id=team.id, gender=Gender.MALE, id_card_number="001095099999",
    )
    make_user(email="chua@company.vn", full_name="Chưa Đăng Ký", employee_code="NV003")

    registration = Registration(
        event_id=event.id, user_id=traveller.id, is_participating=True, shift_id=shift.id,
        status=RegistrationStatus.SUBMITTED, submitted_at=NOW,
    )
    mate_registration = Registration(
        event_id=event.id, user_id=mate.id, is_participating=True,
        status=RegistrationStatus.SUBMITTED, submitted_at=NOW,
    )
    db.add_all([registration, mate_registration])
    db.flush()
    db.add_all(
        [
            RegistrationBusNeed(
                registration_id=registration.id, trip_leg_id=leg.id, needs_bus=True,
                pickup_point_id=pickup.id,
            ),
            RegistrationBusNeed(
                registration_id=mate_registration.id, trip_leg_id=leg.id, needs_bus=True,
                pickup_point_id=pickup.id,
            ),
            FlightAssignment(
                registration_id=registration.id, flight_id=outbound.id,
                direction=FlightDirection.OUTBOUND, assignment_mode=AssignmentMode.AUTO,
                assigned_at=NOW,
            ),
        ]
    )
    bus = Bus(
        event_id=event.id, trip_leg_id=leg.id, bus_code="XE-01", plate_number="29B-123.45",
        capacity=45, pickup_point_id=pickup.id, leader_name="Lê Trưởng Xe", leader_phone="0933000333",
    )
    hotel = Hotel(event_id=event.id, name="Sunset Beach Resort")
    db.add_all([bus, hotel])
    db.flush()
    db.add(
        BusAssignment(
            registration_id=registration.id, bus_id=bus.id, trip_leg_id=leg.id,
            assignment_mode=AssignmentMode.AUTO, assigned_at=NOW,
        )
    )
    room = Room(hotel_id=hotel.id, room_number="1204", capacity=2, gender_policy=RoomGenderPolicy.MALE)
    empty = Room(hotel_id=hotel.id, room_number="1205", capacity=2, gender_policy=RoomGenderPolicy.FEMALE)
    db.add_all([room, empty])
    db.flush()
    db.add_all(
        [
            RoomAssignment(
                registration_id=registration.id, room_id=room.id, is_room_captain=True,
                assignment_mode=AssignmentMode.MANUAL, assigned_at=NOW,
            ),
            RoomAssignment(
                registration_id=mate_registration.id, room_id=room.id,
                assignment_mode=AssignmentMode.MANUAL, assigned_at=NOW,
            ),
        ]
    )
    db.commit()
    return {"event": event.id}


@pytest.fixture
def admin(world, auth_headers):
    return auth_headers("btc@company.vn")


def download(client: TestClient, headers, url: str, **params):
    response = client.get(url, headers=headers, params=params)
    assert response.status_code == 200, response.text
    return response


def sheets(response) -> dict[str, list[dict]]:
    """{tên sheet: [dict theo tiêu đề]}."""
    book = load_workbook(BytesIO(response.content), read_only=True)
    result = {}
    for sheet in book.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        header = rows[0] if rows else ()
        result[sheet.title] = [dict(zip(header, row, strict=False)) for row in rows[1:]]
    book.close()
    return result


def headers_of(response, sheet_name: str) -> list[str]:
    book = load_workbook(BytesIO(response.content), read_only=True)
    header = [cell for cell in next(book[sheet_name].iter_rows(values_only=True)) if cell]
    book.close()
    return header


def export_audits(db: Session) -> list[dict]:
    rows = db.query(AuditLog).filter(AuditLog.action == "export.downloaded").order_by(AuditLog.id).all()
    return [json.loads(row.after_data) for row in rows]


def test_exports_are_organizer_only(client: TestClient, world, auth_headers):
    employee = auth_headers("nv@company.vn")
    for url in EXPORTS:
        assert client.get(url, headers=employee).status_code == 403, url
        assert client.get(url).status_code == 401, url


def test_download_headers_forbid_caching(client: TestClient, admin):
    response = download(client, admin, "/api/v1/rooms/export")
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["cache-control"] == "no-store"
    disposition = response.headers["content-disposition"]
    assert disposition.startswith('attachment; filename="phan-phong-tb2026-')
    assert disposition.endswith(".xlsx")


def test_user_export_adds_sensitive_columns_only_when_asked(client: TestClient, admin, db: Session):
    basic = download(client, admin, "/api/v1/admin/users/export")
    columns = headers_of(basic, "CBNV")
    assert "Số CCCD/Hộ chiếu" not in columns and "Ngày sinh" not in columns
    assert b"001095012345" not in basic.content

    full = download(client, admin, "/api/v1/admin/users/export", include_sensitive="true")
    rows = {row["Mã NV"]: row for row in sheets(full)["CBNV"]}
    assert rows["NV001"]["Số CCCD/Hộ chiếu"] == "001095012345"
    assert rows["NV001"]["SĐT"] == "0911000111"  # chuỗi, không mất số 0
    assert rows["NV001"]["Đăng ký kỳ này"] == "Tham gia"
    assert rows["NV003"]["Đăng ký kỳ này"] == "Chưa đăng ký"
    assert "Dị ứng hải sản" not in json.dumps(rows, ensure_ascii=False, default=str)

    assert [(item["kind"], item["sensitive"]) for item in export_audits(db)] == [
        ("users", False),
        ("users", True),
    ]


def test_user_supplied_formula_is_written_as_plain_text(client: TestClient, admin):
    response = download(client, admin, "/api/v1/admin/users/export")
    with zipfile.ZipFile(BytesIO(response.content)) as archive:
        sheet_xml = archive.read("xl/worksheets/sheet1.xml").decode()
        strings = archive.read("xl/sharedStrings.xml").decode() if "xl/sharedStrings.xml" in archive.namelist() else ""
    assert "<f>" not in sheet_xml
    assert "HYPERLINK" in strings or "HYPERLINK" in sheet_xml


def test_registration_export_lists_bus_needs_and_non_respondents(client: TestClient, admin):
    data = sheets(download(client, admin, "/api/v1/registrations/export"))

    assert list(data) == ["Đăng ký", "Chưa đăng ký"]
    rows = {row["Mã NV"]: row for row in data["Đăng ký"]}
    assert rows["NV001"]["Xe: HN → Sân bay"] == "Có — Toà nhà Keangnam"
    assert rows["NV001"]["Ca nguyện vọng"] == "Ca 1"
    # Người chưa đăng ký có trong sheet thứ hai; tài khoản BTC cũng chưa đăng ký.
    assert {row["Mã NV"] for row in data["Chưa đăng ký"]} == {"NV003", "BTC001"}


def test_flight_manifest_is_logged_as_sensitive(client: TestClient, admin, db: Session):
    data = sheets(download(client, admin, "/api/v1/flights/export"))

    assert list(data) == ["Chiều đi", "Chiều về", "Chưa có chuyến"]
    [passenger] = [row for row in data["Chiều đi"] if row["Mã NV"]]
    assert passenger["Chuyến bay"] == "VN1234"
    assert passenger["Số CCCD/Hộ chiếu"] == "001095012345"
    assert passenger["Khởi hành (giờ VN)"] == "15/10/2026 13:30"
    missing = {(row["Chiều"], row["Mã NV"]) for row in data["Chưa có chuyến"]}
    assert missing == {("Chiều về", "NV001"), ("Chiều đi", "NV002"), ("Chiều về", "NV002")}
    assert export_audits(db) == [{"kind": "flight_manifest", "rows": 1, "sensitive": True}]


def test_bus_export_has_a_sheet_per_leg_with_unassigned_riders(client: TestClient, admin):
    data = sheets(download(client, admin, "/api/v1/buses/export"))

    rows = data["HN → Sân bay"]
    by_code = {row["Mã NV"]: row for row in rows}
    assert by_code["NV001"]["Xe"] == "XE-01"
    assert by_code["NV001"]["Trưởng xe"] == "Lê Trưởng Xe"
    assert by_code["NV002"]["Xe"] == "Chưa có xe"
    assert by_code["NV002"]["Điểm đón đã chọn"] == "Toà nhà Keangnam"


def test_room_export_reimports_unchanged(client: TestClient, admin):
    exported = download(client, admin, "/api/v1/rooms/export")
    data = sheets(exported)
    assert list(data) == ["Phân phòng", "Phòng trống", "Chưa có phòng"]
    assert [row["Trưởng phòng"] for row in data["Phân phòng"]] == ["x", None]
    assert [row["Số phòng"] for row in data["Phòng trống"]] == ["1205"]

    response = client.post(
        "/api/v1/rooms/import",
        headers=admin,
        params={"dry_run": "true"},
        files={"file": ("phan-phong.xlsx", exported.content, exported.headers["content-type"])},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["error_count"] == 0
    assert (body["to_create"], body["to_move"], body["unchanged"]) == (0, 0, 2)


def test_user_export_reimports_unchanged(client: TestClient, admin):
    exported = download(client, admin, "/api/v1/admin/users/export", include_sensitive="true")
    response = client.post(
        "/api/v1/admin/users/import",
        headers=admin,
        files={"file": ("cbnv.xlsx", exported.content, exported.headers["content-type"])},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["error_count"] == 0, body["errors"]
    assert (body["to_create"], body["to_update"], body["unchanged"]) == (0, 0, 4)
